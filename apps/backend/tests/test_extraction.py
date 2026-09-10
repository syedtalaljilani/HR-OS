import base64
import io
import uuid

import pymupdf
import pytest

from app.services import extraction_service

TXT = (
    "Alice Smith\n"
    "Email: alice@example.com\n"
    "Phone: 03001234567\n"
    "Address: Multan\n"
    "Skills: Python, SQL, FastAPI\n"
    "Education: B.S. Computer Science - COMSATS\n"
    "Experience: Developer at XY Corp (2021-present)\n"
)


@pytest.fixture(autouse=True)
def _clear_profile_cache():
    extraction_service._PROFILE_CACHE.clear()
    yield
    extraction_service._PROFILE_CACHE.clear()


def _make_text_pdf(text: str = TXT) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 60), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def _make_image_pdf(text: str = TXT) -> bytes:
    """A PDF with NO text layer (rasterized page) -> forces the OCR path."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 60), text, fontsize=11)
    pix = page.get_pixmap(dpi=150)
    png = pix.tobytes("png")
    doc.close()
    out = pymupdf.open()
    op = out.new_page(width=595, height=842)
    op.insert_image(op.rect, stream=png)
    data = out.tobytes()
    out.close()
    return data


def _fake_chat_json(messages, model=None, max_tokens=None, temperature=0.1):
    return {
        "name": "Alice Smith",
        "email": "alice@example.com",
        "phone": "03001234567",
        "address": "Multan",
        "expected_salary": None,
        "summary": None,
        "skills": ["Python", "SQL", "FastAPI"],
        "languages": [],
        "interests": [],
        "links": [],
        "education": [
            {"degree": "B.S. Computer Science", "institution": "COMSATS", "years": "2016-2020"}
        ],
        "experience": [
            {"position": "Developer", "company": "XY Corp", "years": "2021-present", "description": "built stuff"}
        ],
        "projects": [],
        "certifications": [],
        "publications": [],
    }


@pytest.fixture
def no_models(monkeypatch):
    from app.services import ai_client

    monkeypatch.setattr(ai_client, "chat_json", _fake_chat_json)
    return ai_client


class TestTextExtraction:
    def test_pdf_native_text(self):
        text = extraction_service.extract_text_from_cv_bytes(
            _make_text_pdf(), "cv.pdf"
        )
        assert "Alice Smith" in text
        assert "Python" in text

    def test_txt(self):
        text = extraction_service.extract_text_from_cv_bytes(
            TXT.encode(), "cv.txt"
        )
        assert "alice@example.com" in text

    def test_docx(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Alice Smith")
        doc.add_paragraph("Email: alice@example.com")
        buf = io.BytesIO()
        doc.save(buf)
        text = extraction_service.extract_text_from_cv_bytes(
            buf.getvalue(), "cv.docx"
        )
        assert "Alice Smith" in text

    def test_unsupported_type(self):
        with pytest.raises(Exception) as exc:
            extraction_service.extract_text_from_cv_bytes(
                b"not a real doc", "cv.rtf"
            )
        assert exc.type.__name__ == "HTTPException"


class TestProfileExtraction:
    def test_native_pdf_fast_path(self, no_models):
        no_models.ocr_images = lambda *a, **k: (_ for _ in ()).throw(
            AssertionError("OCR should not run on a text-layer PDF")
        )
        result = extraction_service.extract_profile_from_cv(
            _make_text_pdf(), "cv.pdf"
        )
        assert result["ocr"] is False
        assert result["profile"]["name"] == "Alice Smith"
        assert result["profile"]["email"] == "alice@example.com"
        assert "Python" in result["profile"]["skills"]

    def test_scanned_pdf_uses_ocr(self, no_models, monkeypatch):
        captured = {}

        def fake_ocr(images_b64, model=None):
            captured["pages"] = len(images_b64)
            captured["model"] = model
            return TXT

        monkeypatch.setattr(no_models, "ocr_images", fake_ocr)
        monkeypatch.setattr(
            "app.core.config.settings.OLLAMA_OCR_MODEL", "deepseek-ocr-test"
        )
        result = extraction_service.extract_profile_from_cv(
            _make_image_pdf(), "scanned.pdf"
        )
        assert result["ocr"] is True
        assert captured["pages"] == 1
        assert captured["model"] == "deepseek-ocr-test"
        assert result["profile"]["name"] == "Alice Smith"

    def test_ocr_unavailable_falls_back_to_400(self, no_models, monkeypatch):
        from app.services.ai_client import AIUnavailable

        def down(images_b64, model=None):
            raise AIUnavailable("model down")

        monkeypatch.setattr(no_models, "ocr_images", down)
        with pytest.raises(Exception) as exc:
            extraction_service.extract_profile_from_cv(
                _make_image_pdf(), "scanned.pdf"
            )
        assert exc.type.__name__ == "HTTPException"
        assert "Could not read any text" in str(exc.value)

    def test_ai_unavailable_uses_regex_fallback(self, no_models, monkeypatch):
        monkeypatch.setattr(
            no_models, "chat_json", raise_ai_unavailable
        )
        result = extraction_service.extract_profile_from_cv(
            _make_text_pdf(), "cv.pdf"
        )
        assert result["profile"]["name"] == "Alice Smith"
        assert result["profile"]["email"] == "alice@example.com"

    def test_profile_cache_skips_second_model_call(self, no_models):
        calls = []

        def counting(messages, model=None, max_tokens=None, temperature=0.1):
            calls.append(1)
            return _fake_chat_json(messages)

        no_models.chat_json = counting
        contents = _make_text_pdf()
        extraction_service.extract_profile_from_cv(contents, "cv.pdf")
        extraction_service.extract_profile_from_cv(contents, "cv.pdf")
        assert len(calls) == 1


class TestOcrHelpers:
    def test_render_pdf_pages_returns_base64_pngs(self):
        pages = extraction_service._render_pdf_pages_base64(_make_image_pdf())
        assert len(pages) == 1
        raw = base64.b64decode(pages[0])
        assert raw[:8] == b"\x89PNG\r\n\x1a\n"

    def test_ocr_pages_are_capped(self, no_models, monkeypatch):
        # Build a 3-page image-only PDF; OCR_MAX_PAGES=2 -> only 2 rendered.
        single = _make_image_pdf()
        doc = pymupdf.open()
        for _ in range(3):
            inner = pymupdf.open(stream=single, filetype="pdf")
            doc.insert_pdf(inner)
            inner.close()
        three_page = doc.tobytes()
        doc.close()

        monkeypatch.setattr(
            extraction_service.settings, "OCR_MAX_PAGES", 2
        )
        pages = extraction_service._render_pdf_pages_base64(three_page)
        assert len(pages) == 2

    def test_strip_ocr_markers(self):
        dirty = (
            "### <|end|>\n"
            "SYED TALAL JILANI   \n"
            "Email:   alice@example.com\n"
            "<|im_start|>user\n"
            "<|im_end|>user\n"
            "Skills: Python\n"
        )
        clean = extraction_service._strip_ocr_markers(dirty)
        assert "SYED TALAL JILANI" in clean
        assert "alice@example.com" in clean
        assert "Skills: Python" in clean
        assert "###" not in clean
        assert "<|" not in clean


class TestValidate:
    def test_valid_profile(self):
        profile = _fake_chat_json([])
        report = extraction_service.validate_cv(TXT, profile)
        assert report["valid"] is True
        assert report["requires_review"] is False

    def test_short_text_flags_issue(self):
        report = extraction_service.validate_cv("short", _fake_chat_json([]))
        assert report["valid"] is False
        assert any("too little text" in i for i in report["issues"])


def raise_ai_unavailable(*args, **kwargs):
    from app.services.ai_client import AIUnavailable

    raise AIUnavailable("no model")